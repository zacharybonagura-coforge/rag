from collections.abc import Sequence
from pathlib import Path

import pytest
from docx import Document

from ingest import DocumentStatus, LoadError, _soft_end, _windows, chunk, ingest
from schemas.chunk import Chunk

CORPUS_FILE = Path("data/corpus-tiny/harbor-bike-shop-handbook.md")
CORPUS_DIR = Path("data/corpus-tiny")

class FakeEmbedder:
    provider = "fake"
    model_id = "fake"
    dimension = 768

    def __init__(self, vectors: list[list[float]] | None = None) -> None:
        self.vectors = vectors

    def embed_query(self, query: str) -> list[float]:
        return [0.0] * self.dimension

    def embed_documents(self, texts: Sequence[str]) -> list[list[float]]:
        if self.vectors is not None:
            return self.vectors
        return [[float(i)] + [0.0] * 767 for i, _ in enumerate(texts)]


def test_chunk_handbook_one_section_each() -> None:
    records = chunk(CORPUS_FILE)

    assert [r["section_title"] for r in records] == [
        "Hours",
        "Rentals",
    ]
    assert records[0]["chunk_id"] == "harbor-bike-shop-handbook:v1.0:section-1"
    assert records[1]["chunk_id"] == "harbor-bike-shop-handbook:v1.0:section-2"
    assert records[0]["document"] == "harbor-bike-shop-handbook"
    assert records[0]["version"] == "1.0"
    assert "Tuesday" in str(records[0]["text"])
    assert all(r["section"] == str(i) for i, r in enumerate(records, start=1))


def test_chunk_skips_title() -> None:
    records = chunk(CORPUS_FILE)

    assert not any("Handbook" in str(r["section_title"]) for r in records)
    assert not any("Handbook" in str(r["text"]) for r in records)


def test_ingest_builds_chunks_with_embeddings() -> None:
    chunks, _ = ingest(CORPUS_DIR, FakeEmbedder())

    assert len(chunks) == 2
    assert all(isinstance(c, Chunk) for c in chunks)
    assert chunks[0].chunk_id == "harbor-bike-shop-handbook:v1.0:section-1"
    assert len(chunks[0].embedding) == 768
    assert chunks[1].section_title == "Rentals"


def test_chunk_missing_file_raises(tmp_path: Path) -> None:
    with pytest.raises(FileNotFoundError):
        chunk(tmp_path / "missing.md")


def test_chunk_no_sections_uses_title_as_section(tmp_path: Path) -> None:
    path = tmp_path / "empty.md"
    path.write_text("# Title only (Version 1)\n\nPreamble with no sections.\n")
    records = chunk(path)

    assert len(records) == 1
    assert records[0]["section_title"] == "Title only (Version 1)"
    assert records[0]["chunk_id"] == "empty:v1:section-1"
    assert records[0]["subsection"] == ""
    assert "Preamble with no sections." in str(records[0]["text"])


def test_chunk_plain_body_uses_filename_as_section(tmp_path: Path) -> None:
    path = tmp_path / "memo.md"
    path.write_text("Just a paragraph. No headings at all.\n")
    records = chunk(path)

    assert len(records) == 1
    assert records[0]["section_title"] == "memo"
    assert records[0]["version"] == ""
    assert records[0]["chunk_id"] == "memo:v:section-1"
    assert "Just a paragraph." in str(records[0]["text"])


def test_chunk_plain_body_splits_long_text(tmp_path: Path) -> None:
    path = tmp_path / "notes.md"
    path.write_text(("word " * 80).strip() + "\n")
    records = chunk(path, max_chars=80, overlap=10)

    assert len(records) > 1
    assert [r["chunk_id"] for r in records] == [
        f"notes:v:section-1:part-{i}" for i in range(len(records))
    ]
    assert all(len(str(r["text"])) <= 80 for r in records)


def test_chunk_title_without_version_ok_if_no_sections(tmp_path: Path) -> None:
    path = tmp_path / "loose.md"
    path.write_text("# Harbor Bike Shop Handbook\n\nOpen daily.\n")
    records = chunk(path)

    assert len(records) == 1
    assert records[0]["chunk_id"] == "loose:v:section-1"
    assert records[0]["section_title"] == "Harbor Bike Shop Handbook"
    assert "Open daily." in str(records[0]["text"])


def test_chunk_skips_empty_section(tmp_path: Path) -> None:
    path = tmp_path / "gaps.md"
    path.write_text(
        "# Doc (Version 1)\n\n## Hours\nOpen daily.\n\n## Empty\n\n## Rentals\nBring an ID.\n"
    )

    records = chunk(path)

    assert [r["section_title"] for r in records] == ["Hours", "Rentals"]
    assert records[1]["chunk_id"] == "gaps:v1:section-2"


def test_ingest_raises_when_embedder_count_mismatches() -> None:
    embedder = FakeEmbedder(vectors=[[0.0] * 768])

    with pytest.raises(ValueError, match="shorter than argument 1"):
        ingest(CORPUS_DIR, embedder)


def test_soft_end_prefers_newline_in_last_fifth() -> None:
    window = ("a" * 40) + "line\nmore"
    end = _soft_end(window)

    assert window[end - 1] == "\n"
    assert window[:end].endswith("line\n")


def test_soft_end_prefers_sentence_when_no_newline() -> None:
    window = ("a" * 40) + "Hi. There"
    end = _soft_end(window)

    assert window[:end].endswith("Hi. ")


def test_soft_end_falls_back_to_space() -> None:
    window = ("a" * 40) + "hi there"
    end = _soft_end(window)

    assert window[:end].endswith("hi ")


def test_soft_end_returns_full_window_without_separator() -> None:
    window = "a" * 50

    assert _soft_end(window) == len(window)


def test_windows_keeps_short_text() -> None:
    assert _windows("short", max_chars=10, overlap=2) == ["short"]


def test_windows_parts_stay_within_limit() -> None:
    text = "word " * 80
    parts = _windows(text, max_chars=40, overlap=8)

    assert len(parts) > 1
    assert all(len(part) <= 40 for part in parts)


def test_windows_overlap_starts_on_a_word() -> None:
    text = "alpha beta gamma delta epsilon zeta eta theta"
    parts = _windows(text, max_chars=20, overlap=8)

    assert len(parts) > 1
    assert all(not part[0].isspace() for part in parts)
    assert " " + parts[1].split()[0] in text


def test_windows_keeps_table_rows_whole() -> None:
    text = (
        ("word " * 60).strip()
        + "\nIssue | Fee\nLate return | One extra day\nLost lock | $20"
    )
    parts = _windows(text, max_chars=80, overlap=20)

    assert any("Lost lock | $20" in part for part in parts)
    assert all("Lost lock |" not in part or "Lost lock | $20" in part for part in parts)


def test_chunk_markdown_subsections(tmp_path: Path) -> None:
    path = tmp_path / "hours.md"
    path.write_text(
        "# Doc (Version 1)\n\n"
        "## Hours\nIntro line.\n\n"
        "### Weekday\nOpen at 10.\n\n"
        "### Weekend\nClosed.\n"
    )
    records = chunk(path)

    assert [r["chunk_id"] for r in records] == [
        "hours:v1:section-1:sub-0",
        "hours:v1:section-1:sub-1",
        "hours:v1:section-1:sub-2",
    ]
    assert [r["subsection_title"] for r in records] == ["", "Weekday", "Weekend"]
    assert records[0]["text"] == "Intro line."
    assert records[1]["text"] == "Open at 10."


def test_chunk_markdown_splits_long_section(tmp_path: Path) -> None:
    path = tmp_path / "notes.md"
    path.write_text("# Doc (Version 1)\n\n## Notes\n" + ("word " * 80) + "\n")
    records = chunk(path, max_chars=80, overlap=10)

    assert all(r["section_title"] == "Notes" for r in records)
    assert [r["chunk_id"] for r in records] == [
        f"notes:v1:section-1:part-{i}" for i in range(len(records))
    ]
    assert len(records) > 1
    assert all(len(str(r["text"])) <= 80 for r in records)


def test_chunk_markdown_title_without_version_raises(tmp_path: Path) -> None:
    path = tmp_path / "bad.md"
    path.write_text("# Harbor Bike Shop Handbook\n\n## Hours\nOpen daily.\n")
    
    with pytest.raises(ValueError, match="has sections but no version"):
        chunk(path)


def _write_docx(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("Harbor Bike Shop Handbook (Version 1.0)", style="Heading 1")
    doc.add_paragraph("Hours", style="Heading 2")
    doc.add_paragraph("Walk-ins are welcome.")
    doc.add_paragraph("Weekday", style="Heading 3")
    doc.add_paragraph("Open from 10am to 6pm.")
    doc.add_paragraph("Prices", style="Heading 2")
    doc.add_paragraph("See the table.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "City"
    table.cell(0, 1).text = "$25"
    table.cell(1, 0).text = "Electric"
    table.cell(1, 1).text = "$45"
    doc.add_paragraph("A photo ID is required.")
    doc.save(str(path))


def test_chunk_docx_subsections_and_table(tmp_path: Path) -> None:
    path = tmp_path / "handbook.docx"
    _write_docx(path)
    records = chunk(path)

    assert [r["chunk_id"] for r in records] == [
        "handbook:v1.0:section-1:sub-0",
        "handbook:v1.0:section-1:sub-1",
        "handbook:v1.0:section-2",
    ]
    assert records[0]["text"] == "Walk-ins are welcome."
    assert records[1]["subsection_title"] == "Weekday"
    assert str(records[2]["text"]) == (
        "See the table.\nCity | $25\nElectric | $45\nA photo ID is required."
    )


def test_ingest_reads_docx(tmp_path: Path) -> None:
    _write_docx(tmp_path / "handbook.docx")
    chunks, _ = ingest(tmp_path, FakeEmbedder())

    assert [c.chunk_id for c in chunks] == [
        "handbook:v1.0:section-1:sub-0",
        "handbook:v1.0:section-1:sub-1",
        "handbook:v1.0:section-2",
    ]
    assert "City | $25" in chunks[2].text


def _pdf_pages() -> list[dict[str, object]]:
    return [
        {
            "metadata": {"page": 1},
            "text": (
                "# Girls' House Rules (Version 1.0)\n\n"
                "## Kitchen\n"
                "The girls eat first.\n\n"
                "### Cookies\n"
                "Leave three cookies for Agnes.\n"
                "|Item|Count|\n"
                "|---|---|\n"
                "|Cookies|3|\n"
            ),
        },
        {
            "metadata": {"page_number": 2},
            "text": (
                "## Emergencies\n"
                "Take the girls to the orange couch.\n"
            ),
        },
    ]


def test_chunk_pdf_subsections_and_table(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    path = tmp_path / "handbook.pdf"
    path.write_bytes(b"%PDF")
    monkeypatch.setattr("blocks.pymupdf4llm.to_markdown", lambda *_a, **_k: _pdf_pages())
    records = chunk(path)

    assert [r["chunk_id"] for r in records] == [
        "handbook:v1.0:section-1:sub-0",
        "handbook:v1.0:section-1:sub-1",
        "handbook:v1.0:section-2",
    ]
    assert records[0]["page"] == 1
    assert records[0]["text"] == "The girls eat first."
    assert records[1]["subsection_title"] == "Cookies"
    assert "Cookies | 3" in str(records[1]["text"])
    assert records[2]["page"] == 2
    assert records[2]["section_title"] == "Emergencies"


def test_ingest_reads_pdf(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    path = tmp_path / "handbook.pdf"
    path.write_bytes(b"%PDF")
    monkeypatch.setattr("blocks.pymupdf4llm.to_markdown", lambda *_a, **_k: _pdf_pages())
    chunks, _ = ingest(tmp_path, FakeEmbedder())

    assert [c.chunk_id for c in chunks] == [
        "handbook:v1.0:section-1:sub-0",
        "handbook:v1.0:section-1:sub-1",
        "handbook:v1.0:section-2",
    ]
    assert chunks[1].text.count("Cookies") >= 1
    assert chunks[2].page == 2


def test_chunk_empty_file_returns_no_records(tmp_path: Path) -> None:
    path = tmp_path / "blank.md"
    path.write_text("")
    assert chunk(path) == []


def test_chunk_title_only_returns_no_records(tmp_path: Path) -> None:
    path = tmp_path / "title.md"
    path.write_text("# Title only (Version 1)\n")
    assert chunk(path) == []


def test_ingest_flags_no_text_and_keeps_good_file(tmp_path: Path) -> None:
    (tmp_path / "blank.md").write_text("")
    (tmp_path / "ok.md").write_text("# Ok (Version 1.0)\n\n## Hours\nOpen daily.\n")
    chunks, documents = ingest(tmp_path, FakeEmbedder())
    assert [c.document for c in chunks] == ["ok"]
    assert documents == [
        DocumentStatus("blank", False, "no_text"),
        DocumentStatus("ok", True, ""),
    ]


def test_ingest_flags_unparsable(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    (tmp_path / "bad.md").write_text("# Bad (Version 1.0)\n\n## Hours\nOpen.\n")
    def boom(path: Path) -> list[object]:
        raise ValueError("corrupt")
    monkeypatch.setattr("ingest.load", boom)
    chunks, documents = ingest(tmp_path, FakeEmbedder())
    assert chunks == []
    assert documents == [DocumentStatus("bad", False, "unparsable")]


def test_chunk_load_error_wraps(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    path = tmp_path / "bad.md"
    path.write_text("x")
    monkeypatch.setattr("ingest.load", lambda _path: (_ for _ in ()).throw(ValueError("corrupt")))
    with pytest.raises(LoadError, match="unparsable"):
        chunk(path)


def test_windows_advances_one_char_when_overlap_covers_window() -> None:
    parts = _windows("a bcdefghijkl", max_chars=5, overlap=4)
    assert len(parts) > 1


def test_ingest_raises_when_directory_has_no_handbooks(tmp_path: Path) -> None:
    with pytest.raises(FileNotFoundError, match="no markdown"):
        ingest(tmp_path, FakeEmbedder())


def test_windows_advances_one_char_when_overlap_covers_window() -> None:
    text = "xxxxxxxx x" + "y" * 20
    parts = _windows(text, max_chars=10, overlap=9)
    assert len(parts) > 1
