from pathlib import Path

import pytest
from docx import Document

from blocks import load


def _write_docx(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("Harbor Bike Shop Handbook (Version 1.0)", style="Heading 1")
    doc.add_paragraph("Hours", style="Heading 2")
    doc.add_paragraph("Walk-ins are welcome.")
    doc.add_paragraph("Weekday", style="Heading 3")
    doc.add_paragraph("Open from 10am to 6pm.")
    doc.add_paragraph("Prices", style="Heading 2")
    doc.add_paragraph("See the table.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "City"
    table.cell(0, 1).text = "$25"
    table.cell(1, 0).text = "Electric"
    table.cell(1, 1).text = "$45"
    doc.add_paragraph("A photo ID is required.")
    doc.save(str(path))


def _write_md(path: Path, text: str) -> Path:
    path.write_text(text)
    return path


def _pdf(tmp_path: Path) -> Path:
    path = tmp_path / "handbook.pdf"
    path.write_bytes(b"%PDF")
    return path


def _stub_pdf(monkeypatch: pytest.MonkeyPatch, text: str) -> None:
    monkeypatch.setattr(
        "blocks.pymupdf4llm.to_markdown",
        lambda *_a, **_k: [{"metadata": {"page": 1}, "text": text}],
    )


def _pdf_pages() -> list[dict[str, object]]:
    return [
        {
            "metadata": {"page": 1},
            "text": (
                "# Girls' House Rules (Version 1.0)\n\n"
                "## Kitchen\n"
                "The girls eat first.\n\n"
                "### Cookies\n"
                "Leave three cookies for Agnes.\n"
                "|Item|Count|\n"
                "|---|---|\n"
                "|Cookies|3|\n"
            ),
        },
        {
            "metadata": {"page_number": 2},
            "text": (
                "## Emergencies\n"
                "Take the girls to the orange couch.\n"
            ),
        },
    ]


def test_load_docx_headings_and_table(tmp_path: Path) -> None:
    path = tmp_path / "handbook.docx"
    _write_docx(path)

    assert [(b.kind, b.text) for b in load(path)] == [
        ("title", "Harbor Bike Shop Handbook (Version 1.0)"),
        ("section", "Hours"),
        ("body", "Walk-ins are welcome."),
        ("subsection", "Weekday"),
        ("body", "Open from 10am to 6pm."),
        ("section", "Prices"),
        ("body", "See the table."),
        ("body", "City | $25"),
        ("body", "Electric | $45"),
        ("body", "A photo ID is required."),
    ]


def test_load_rejects_unknown_suffix(tmp_path: Path) -> None:
    path = tmp_path / "notes.txt"
    path.write_text("hello")

    with pytest.raises(ValueError, match="unsupported file type"):
        load(path)


def test_load_pdf_headings_table_and_page(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(
        "blocks.pymupdf4llm.to_markdown", lambda *_a, **_k: _pdf_pages()
    )

    blocks = [(b.kind, b.text, b.page) for b in load(_pdf(tmp_path)) if b.text]

    assert blocks == [
        ("title", "Girls' House Rules (Version 1.0)", 1),
        ("section", "Kitchen", 1),
        ("body", "The girls eat first.", 1),
        ("subsection", "Cookies", 1),
        ("body", "Leave three cookies for Agnes.", 1),
        ("body", "Item | Count", 1),
        ("body", "Cookies | 3", 1),
        ("section", "Emergencies", 2),
        ("body", "Take the girls to the orange couch.", 2),
    ]


def test_load_pdf_requires_page_chunks_list(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(
        "blocks.pymupdf4llm.to_markdown", lambda *_a, **_k: "# not a list"
    )

    with pytest.raises(TypeError, match="page_chunks"):
        load(_pdf(tmp_path))


def test_load_markdown_flattens_table_and_drops_images(tmp_path: Path) -> None:
    path = _write_md(
        tmp_path / "prices.md",
        "# Doc (Version 1.0)\n\n"
        "## Prices\n"
        "See the table.\n"
        "|Item|Count|\n"
        "| :--- | ---: |\n"
        "|Cookies|3|\n"
        "![cookies](cookies.png)\n"
        "Leave three ![note](n.png) on the tray.\n",
    )

    assert [(b.kind, b.text) for b in load(path)] == [
        ("title", "Doc (Version 1.0)"),
        ("section", "Prices"),
        ("body", "See the table."),
        ("body", "Item | Count"),
        ("body", "Cookies | 3"),
        ("body", "Leave three  on the tray."),
    ]


def test_load_markdown_skips_separator_and_empty_table_rows(tmp_path: Path) -> None:
    path = _write_md(
        tmp_path / "empty.md",
        "# Doc (Version 1.0)\n\n"
        "## Prices\n"
        "|---|---|\n"
        "| | |\n",
    )

    assert [(b.kind, b.text) for b in load(path)] == [
        ("title", "Doc (Version 1.0)"),
        ("section", "Prices"),
    ]


def test_load_markdown_leaves_prose_pipes_and_fenced_tables(tmp_path: Path) -> None:
    path = _write_md(
        tmp_path / "notes.md",
        "# Doc (Version 1.0)\n\n"
        "## Notes\n"
        "Hours | 10am\n"
        "```\n"
        "|Item|Count|\n"
        "|---|---|\n"
        "|Cookies|3|\n"
        "```\n"
        "![alt][ref]\n",
    )

    assert [(b.kind, b.text) for b in load(path)] == [
        ("title", "Doc (Version 1.0)"),
        ("section", "Notes"),
        ("body", "Hours | 10am"),
        ("body", "```"),
        ("body", "|Item|Count|"),
        ("body", "|---|---|"),
        ("body", "|Cookies|3|"),
        ("body", "```"),
        ("body", "![alt][ref]"),
    ]


def test_load_pdf_flattens_table_and_drops_images(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _stub_pdf(
        monkeypatch,
        "# Doc (Version 1.0)\n\n"
        "## Prices\n"
        "See the table.\n"
        "|Item|Count|\n"
        "| :--- | ---: |\n"
        "|Cookies|3|\n"
        "![cookies](cookies.png)\n"
        "Leave three ![note](n.png) on the tray.\n",
    )

    assert [(b.kind, b.text, b.page) for b in load(_pdf(tmp_path))] == [
        ("title", "Doc (Version 1.0)", 1),
        ("section", "Prices", 1),
        ("body", "See the table.", 1),
        ("body", "Item | Count", 1),
        ("body", "Cookies | 3", 1),
        ("body", "Leave three  on the tray.", 1),
    ]


def test_load_pdf_skips_separator_and_empty_table_rows(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _stub_pdf(
        monkeypatch,
        "# Doc (Version 1.0)\n\n## Prices\n|---|---|\n| | |\n",
    )

    assert [(b.kind, b.text) for b in load(_pdf(tmp_path))] == [
        ("title", "Doc (Version 1.0)"),
        ("section", "Prices"),
    ]


def test_load_pdf_leaves_prose_pipes_and_fenced_tables(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _stub_pdf(
        monkeypatch,
        "# Doc (Version 1.0)\n\n"
        "## Notes\n"
        "Hours | 10am\n"
        "```\n"
        "|Item|Count|\n"
        "|---|---|\n"
        "|Cookies|3|\n"
        "```\n"
        "![alt][ref]\n",
    )

    assert [(b.kind, b.text) for b in load(_pdf(tmp_path))] == [
        ("title", "Doc (Version 1.0)"),
        ("section", "Notes"),
        ("body", "Hours | 10am"),
        ("body", "```"),
        ("body", "|Item|Count|"),
        ("body", "|---|---|"),
        ("body", "|Cookies|3|"),
        ("body", "```"),
        ("body", "![alt][ref]"),
    ]


def test_load_pdf_drops_non_int_page(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(
        "blocks.pymupdf4llm.to_markdown",
        lambda *_a, **_k: [{"metadata": {"page": "1"}, "text": "# Doc (Version 1.0)\n"}],
    )
    blocks = load(_pdf(tmp_path))
    assert blocks[0].page is None
